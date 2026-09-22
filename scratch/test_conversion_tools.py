import io
import sys
from pathlib import Path
import pymupdf as fitz
import docx
import openpyxl
import pptx
from fastapi.testclient import TestClient

backend_dir = Path(__file__).parent.parent / "backend"
sys.path.insert(0, str(backend_dir))

from main import app

def create_sample_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text(fitz.Point(100, 100), "Sample Title Header", fontsize=20)
    page.insert_text(fitz.Point(100, 140), "This is a test paragraph content inside the PDF.", fontsize=12)
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()

def create_sample_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Sample Word Document", level=1)
    doc.add_paragraph("This is sample paragraph content in Word.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def create_sample_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Report"
    ws.append(["Product", "Q1 Sales", "Q2 Sales"])
    ws.append(["Item A", 1500, 2300])
    ws.append(["Item B", 800, 1200])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def create_sample_pptx() -> bytes:
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Sample PowerPoint Slide"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()

def test_conversions():
    client = TestClient(app)
    pdf_bytes = create_sample_pdf()
    docx_bytes = create_sample_docx()
    xlsx_bytes = create_sample_xlsx()
    pptx_bytes = create_sample_pptx()

    print("Testing /api/pdf-to-word...")
    res = client.post("/api/pdf-to-word", files={"file": ("test.pdf", pdf_bytes, "application/pdf")})
    assert res.status_code == 200, res.text
    assert len(res.content) > 100
    print("  [OK] /api/pdf-to-word OK")

    print("Testing /api/pdf-to-excel...")
    res = client.post("/api/pdf-to-excel", files={"file": ("test.pdf", pdf_bytes, "application/pdf")})
    assert res.status_code == 200, res.text
    assert len(res.content) > 100
    print("  [OK] /api/pdf-to-excel OK")

    print("Testing /api/pdf-to-pptx...")
    res = client.post("/api/pdf-to-pptx", files={"file": ("test.pdf", pdf_bytes, "application/pdf")})
    assert res.status_code == 200, res.text
    assert len(res.content) > 100
    print("  [OK] /api/pdf-to-pptx OK")

    print("Testing /api/pdf-to-png...")
    res = client.post("/api/pdf-to-png", files={"file": ("test.pdf", pdf_bytes, "application/pdf")})
    assert res.status_code == 200, res.text
    assert len(res.content) > 100
    print("  [OK] /api/pdf-to-png OK")

    print("Testing /api/pdf-to-txt...")
    res = client.post("/api/pdf-to-txt", files={"file": ("test.pdf", pdf_bytes, "application/pdf")})
    assert res.status_code == 200, res.text
    assert b"Sample Title Header" in res.content
    print("  [OK] /api/pdf-to-txt OK")

    print("Testing /api/pdf-to-html...")
    res = client.post("/api/pdf-to-html", files={"file": ("test.pdf", pdf_bytes, "application/pdf")})
    assert res.status_code == 200, res.text
    assert b"<!DOCTYPE html>" in res.content
    print("  [OK] /api/pdf-to-html OK")

    print("Testing /api/pdf-to-pdfa...")
    res = client.post("/api/pdf-to-pdfa", files={"file": ("test.pdf", pdf_bytes, "application/pdf")})
    assert res.status_code == 200, res.text
    out_doc = fitz.open(stream=res.content, filetype="pdf")
    assert out_doc.page_count == 1
    out_doc.close()
    print("  [OK] /api/pdf-to-pdfa OK")

    print("Testing /api/word-to-pdf...")
    res = client.post("/api/word-to-pdf", files={"file": ("test.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert res.status_code == 200, res.text
    out_doc = fitz.open(stream=res.content, filetype="pdf")
    assert out_doc.page_count >= 1
    out_doc.close()
    print("  [OK] /api/word-to-pdf OK")

    print("Testing /api/excel-to-pdf...")
    res = client.post("/api/excel-to-pdf", files={"file": ("test.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert res.status_code == 200, res.text
    out_doc = fitz.open(stream=res.content, filetype="pdf")
    assert out_doc.page_count >= 1
    out_doc.close()
    print("  [OK] /api/excel-to-pdf OK")

    print("Testing /api/pptx-to-pdf...")
    res = client.post("/api/pptx-to-pdf", files={"file": ("test.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
    assert res.status_code == 200, res.text
    out_doc = fitz.open(stream=res.content, filetype="pdf")
    assert out_doc.page_count >= 1
    out_doc.close()
    print("  [OK] /api/pptx-to-pdf OK")

    print("Testing /api/html-to-pdf...")
    res = client.post("/api/html-to-pdf", data={"html": "<h1>Hello HTML</h1><p>Sample paragraph</p>"})
    assert res.status_code == 200, res.text
    out_doc = fitz.open(stream=res.content, filetype="pdf")
    assert out_doc.page_count >= 1
    out_doc.close()
    print("  [OK] /api/html-to-pdf OK")

    print("\nALL CONVERSION TOOL TESTS PASSED SUCCESSFULLY!")

if __name__ == "__main__":
    test_conversions()
