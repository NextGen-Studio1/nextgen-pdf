import sys
import os
import io
from PIL import Image
from docx import Document
import openpyxl
from pptx import Presentation

sys.path.insert(0, os.path.join(os.getcwd(), "backend"))

from fastapi.testclient import TestClient
from app.main import app
from unittest.mock import patch
from app.services.rate_limit_service import RateLimitService

client = TestClient(app)

# Dummy PDF
dummy_pdf = b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj 2 0 obj<</Type/Pages/Count 2/Kids[3 0 R 4 0 R]>>endobj 3 0 obj<</Type/Page/MediaBox[0 0 612 792]/Parent 2 0 R>>endobj 4 0 obj<</Type/Page/MediaBox[0 0 612 792]/Parent 2 0 R>>endobj\nxref\n0 5\n0000000000 65535 f\n0000000009 00000 n\n0000000052 00000 n\n0000000115 00000 n\n0000000185 00000 n\ntrailer<</Size 5/Root 1 0 R>>\nstartxref\n255\n%%EOF"

# PNG
img = Image.new('RGB', (100, 100), color='red')
img_buf = io.BytesIO()
img.save(img_buf, format='PNG')
valid_png = img_buf.getvalue()

# DOCX
doc = Document()
doc.add_paragraph('Test content')
doc_buf = io.BytesIO()
doc.save(doc_buf)
valid_docx = doc_buf.getvalue()

# XLSX
wb = openpyxl.Workbook()
ws = wb.active
ws['A1'] = "Test Cell"
xlsx_buf = io.BytesIO()
wb.save(xlsx_buf)
valid_xlsx = xlsx_buf.getvalue()

# PPTX
prs = Presentation()
prs.slides.add_slide(prs.slide_layouts[0])
pptx_buf = io.BytesIO()
prs.save(pptx_buf)
valid_pptx = pptx_buf.getvalue()

tool_tests = [
    ("Compress PDF", "/api/v1/compress", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"level": "balanced"}),
    ("Merge PDF", "/api/v1/merge", [("files", ("t1.pdf", dummy_pdf, "application/pdf")), ("files", ("t2.pdf", dummy_pdf, "application/pdf"))], {}),
    ("Split PDF", "/api/v1/split", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"mode": "every", "ranges": ""}),
    ("Rotate PDF", "/api/v1/rotate", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"angle": "90", "pages": "all"}),
    ("JPG to PDF", "/api/v1/jpg-to-pdf", [("files", ("test.png", valid_png, "image/png"))], {"layout": "portrait", "fit": "fit"}),
    ("Word to PDF", "/api/v1/word-to-pdf", {"file": ("test.docx", valid_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}, {}),
    ("PDF to Word", "/api/v1/pdf-to-word", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
    ("Excel to PDF", "/api/v1/excel-to-pdf", {"file": ("test.xlsx", valid_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, {}),
    ("PDF to Excel", "/api/v1/pdf-to-excel", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
    ("PPTX to PDF", "/api/v1/pptx-to-pdf", {"file": ("test.pptx", valid_pptx, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}, {}),
    ("PDF to PPTX", "/api/v1/pdf-to-pptx", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
    ("PDF to TXT", "/api/v1/pdf-to-txt", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
    ("PDF to Images", "/api/v1/pdf-to-jpg", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"format": "jpg", "ranges": ""}),
    ("PDF to PNG", "/api/v1/pdf-to-png", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"ranges": ""}),
    ("Protect PDF", "/api/v1/protect-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"password": "123"}),
    ("Unlock PDF", "/api/v1/unlock-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"password": ""}),
    ("Organize PDF", "/api/v1/organize", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"order": "2,1"}),
    ("Extract Pages", "/api/v1/extract-pages", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"pages": "1"}),
    ("Delete Pages", "/api/v1/delete-pages", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"pages": "1"}),
    ("Add Image", "/api/v1/add-image", {"file": ("test.pdf", dummy_pdf, "application/pdf"), "image_file": ("stamp.png", valid_png, "image/png")}, {"position": "bottom-right", "width": "120", "height": "120", "pages": "all"}),
    ("Crop PDF", "/api/v1/crop", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"margin": "36", "pages": "all"}),
    ("Watermark PDF", "/api/v1/watermark", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"text": "WATERMARK"}),
    ("Page Numbers", "/api/v1/page-numbers", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"position": "bottom-center"}),
    ("Metadata Editor", "/api/v1/metadata-editor", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"title": "New Title"}),
    ("Remove Metadata", "/api/v1/remove-metadata", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
    ("Flatten PDF", "/api/v1/flatten-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
    ("Repair PDF", "/api/v1/repair-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
    ("Extract Images", "/api/v1/extract-images", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
    ("AI Summarize", "/api/v1/ai-summarize", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"summary_type": "executive"}),
    ("Chat PDF", "/api/v1/chat-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"question": "What is in this PDF?"}),
    ("Translate PDF", "/api/v1/translate-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"target_lang": "es"}),
    ("PDF Bookmarks", "/api/v1/pdf-bookmarks", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"action": "view"}),
    ("Redact PDF", "/api/v1/redact-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"text": "secret"}),
    ("OCR PDF", "/api/v1/ocr-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"language": "eng"}),
    ("HTML to PDF", "/api/v1/html-to-pdf", {}, {"html": "<h1>Test HTML</h1>"}),
    ("PDF to Markdown", "/api/v1/pdf-to-markdown", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
    ("PDF to PDF/A", "/api/v1/pdf-to-pdfa", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
    ("Compare PDFs", "/api/v1/compare-pdfs", {"file1": ("t1.pdf", dummy_pdf, "application/pdf"), "file2": ("t2.pdf", dummy_pdf, "application/pdf")}, {}),
    ("Overlay PDF", "/api/v1/overlay-pdf", {"file": ("t1.pdf", dummy_pdf, "application/pdf"), "overlay_file": ("t2.pdf", dummy_pdf, "application/pdf")}, {}),
    ("Deskew PDF", "/api/v1/deskew-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
    ("Add Links", "/api/v1/add-links", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"url": "https://example.com"}),
    ("Annotate PDF", "/api/v1/annotate-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"comment": "Note text"}),
    ("Highlight PDF", "/api/v1/highlight-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"text": "highlight"}),
    ("Edit PDF", "/api/v1/edit-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"text": "edit text"}),
    ("Resize Pages", "/api/v1/resize-pages", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"paper_size": "a4"}),
    ("Scanned to Text", "/api/v1/scanned-pdf-to-text", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
    ("PDF to HTML", "/api/v1/pdf-to-html", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
    ("Images to PDF", "/api/v1/images-to-pdf", [("files", ("test.png", valid_png, "image/png"))], {"layout": "portrait"})
]

results = []
with patch.object(RateLimitService, 'check_and_increment_quota', return_value=None):
    for idx, (name, ep, files, data) in enumerate(tool_tests, 1):
        try:
            res = client.post(ep, files=files, data=data)
            status = "PASS" if res.status_code == 200 and len(res.content) > 0 else f"FAIL ({res.status_code})"
            results.append((idx, name, ep, res.status_code, len(res.content), status))
            print(f"{idx:2d}. {name:<20} | {ep:<25} | Code: {res.status_code} | Bytes: {len(res.content):>6} | {status}")
        except Exception as e:
            results.append((idx, name, ep, 500, 0, f"FAIL (EXC: {e})"))
            print(f"{idx:2d}. {name:<20} | {ep:<25} | EXC: {e}")

all_pass = all(r[5] == "PASS" for r in results)
print(f"\n==================================================")
print(f"48-TOOL MATRIX VERIFICATION: {len(results)}/48 TOOLS TESTED")
print(f"OVERALL STATUS: {'ALL 48 TOOLS PASS (ALL 200 OK)' if all_pass else 'FAIL'}")
print(f"==================================================")
