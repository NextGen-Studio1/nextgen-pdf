import io
import sys
import os
import json
import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, os.path.join(r"d:\Download\NextGen-PDF-Launch-Ready-V1\nextgen", "backend"))

from app.main import app

client = TestClient(app)

def create_dummy_pdf(pages_count=1):
    import pymupdf as fitz
    doc = fitz.open()
    for p in range(pages_count):
        page = doc.new_page()
        page.insert_text((72, 72), f"NextGen PDF Test Document Page {p+1}", fontsize=18)
    out = io.BytesIO()
    doc.save(out)
    doc.close()
    return out.getvalue()

def create_dummy_img():
    from PIL import Image
    img = Image.new('RGB', (100, 100), color='blue')
    out = io.BytesIO()
    img.save(out, format='PNG')
    return out.getvalue()

def create_dummy_docx():
    import docx
    doc = docx.Document()
    doc.add_heading("Sample Word Document", level=1)
    doc.add_paragraph("This is a valid DOCX document for NextGen PDF conversion testing.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def run_v1_matrix_test():
    dummy_pdf = create_dummy_pdf(1)
    dummy_pdf_2p = create_dummy_pdf(2)
    dummy_img = create_dummy_img()
    dummy_docx = create_dummy_docx()

    tools_spec = [
        ("Merge PDF", "merge", "/api/v1/merge", {"files": [("files", ("test1.pdf", dummy_pdf, "application/pdf")), ("files", ("test2.pdf", dummy_pdf, "application/pdf"))]}, {}),
        ("Compress PDF", "compress", "/api/v1/compress", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"level": "balanced"}),
        ("Rotate PDF", "rotate", "/api/v1/rotate", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"angle": "90", "pages": "all"}),
        ("Protect PDF", "protect", "/api/v1/protect", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"password": "testpass123"}),
        ("Images to PDF", "images-to-pdf", "/api/v1/images-to-pdf", {"files": [("files", ("img1.jpg", dummy_img, "image/jpeg"))]}, {"layout": "portrait", "fit": "fit"}),
        ("Split PDF", "split", "/api/v1/split", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"mode": "every", "ranges": ""}),
        ("Watermark PDF", "watermark", "/api/v1/watermark", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"text": "SAMPLE", "position": "diagonal"}),
        ("Organize PDF", "organize", "/api/v1/organize", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"order": "1"}),
        ("Delete Pages", "delete-pages", "/api/v1/delete-pages", {"file": ("test.pdf", dummy_pdf_2p, "application/pdf")}, {"pages": "1"}),
        ("Extract Pages", "extract-pages", "/api/v1/extract-pages", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"pages": "1"}),
        ("Page Numbers", "page-numbers", "/api/v1/page-numbers", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"position": "bottom-center"}),
        ("Crop PDF", "crop", "/api/v1/crop", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"margin": "36"}),
        ("PDF to Word", "pdf-to-word", "/api/v1/pdf-to-word", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("PDF to Excel", "pdf-to-excel", "/api/v1/pdf-to-excel", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("PDF to PPTX", "pdf-to-pptx", "/api/v1/pdf-to-pptx", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("PDF to PNG", "pdf-to-png", "/api/v1/pdf-to-png", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("PDF to TXT", "pdf-to-txt", "/api/v1/pdf-to-txt", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("PDF to HTML", "pdf-to-html", "/api/v1/pdf-to-html", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("PDF to PDF/A", "pdf-to-pdfa", "/api/v1/pdf-to-pdfa", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("Word to PDF", "word-to-pdf", "/api/v1/word-to-pdf", {"file": ("test.docx", dummy_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}, {}),
        ("Excel to PDF", "excel-to-pdf", "/api/v1/excel-to-pdf", {"file": ("test.xlsx", b"dummy xlsx bytes", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, {}),
        ("PPTX to PDF", "pptx-to-pdf", "/api/v1/pptx-to-pdf", {"file": ("test.pptx", b"dummy pptx bytes", "application/vnd.openxmlformats-officedocument.presentationml.presentation")}, {}),
        ("HTML to PDF", "html-to-pdf", "/api/v1/html-to-pdf", {}, {"html": "<h1>Test Document</h1>"}),
        ("Unlock PDF", "unlock-pdf", "/api/v1/unlock-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"password": ""}),
        ("Encrypt PDF", "encrypt-pdf", "/api/v1/encrypt-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"user_password": "pass123", "owner_password": "pass123"}),
        ("Redact PDF", "redact-pdf", "/api/v1/redact-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"text": "Test"}),
        ("Remove Metadata", "remove-metadata", "/api/v1/remove-metadata", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("Edit PDF", "edit-pdf", "/api/v1/edit-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"text": "Inserted Text"}),
        ("Add Image", "add-image", "/api/v1/add-image", {"file": ("test.pdf", dummy_pdf, "application/pdf"), "image_file": ("img1.jpg", dummy_img, "image/jpeg")}, {"position": "bottom-right"}),
        ("Add Links", "add-links", "/api/v1/add-links", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"url": "https://example.com"}),
        ("Highlight PDF", "highlight-pdf", "/api/v1/highlight-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"text": "NextGen"}),
        ("Annotate PDF", "annotate-pdf", "/api/v1/annotate-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"comment": "Sticky note"}),
        ("Resize Pages", "resize-pages", "/api/v1/resize-pages", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"paper_size": "a4"}),
        ("OCR PDF", "ocr", "/api/v1/ocr-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("Scanned PDF Text", "scanned-pdf-to-text", "/api/v1/scanned-pdf-to-text", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("AI Summarize", "ai-summarize", "/api/v1/ai-summarize", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("Chat PDF", "chat-pdf", "/api/v1/chat-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"question": "What is this?"}),
        ("Translate PDF", "translate-pdf", "/api/v1/translate-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"target_lang": "es"}),
        ("PDF to Markdown", "pdf-to-markdown", "/api/v1/pdf-to-markdown", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("Repair PDF", "repair-pdf", "/api/v1/repair-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("Compare PDFs", "compare-pdfs", "/api/v1/compare-pdfs", {"file1": ("test1.pdf", dummy_pdf, "application/pdf"), "file2": ("test2.pdf", dummy_pdf, "application/pdf")}, {}),
        ("Flatten PDF", "flatten-pdf", "/api/v1/flatten-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("Deskew PDF", "deskew-pdf", "/api/v1/deskew-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("Extract Images", "extract-images", "/api/v1/extract-images", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {}),
        ("Overlay PDF", "overlay-pdf", "/api/v1/overlay-pdf", {"file": ("test.pdf", dummy_pdf, "application/pdf"), "overlay_file": ("overlay.pdf", dummy_pdf, "application/pdf")}, {}),
        ("PDF Bookmarks", "pdf-bookmarks", "/api/v1/pdf-bookmarks", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"action": "view"}),
        ("Metadata Editor", "metadata-editor", "/api/v1/metadata-editor", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {"title": "New Title"}),
        ("PDF to JPG", "pdf-to-images", "/api/v1/pdf-to-jpg", {"file": ("test.pdf", dummy_pdf, "application/pdf")}, {})
    ]

    results = []
    print(f"\n==================================================")
    print(f"RUNNING ALL 48 ROUTE VERIFICATION MATRIX ON /api/v1")
    print(f"==================================================\n")

    routes_registered = [r.path for r in app.routes]

    for idx, (name, tool_key, route, files_payload, data_payload) in enumerate(tools_spec, 1):
        old_route = f"/api/{tool_key}"
        is_registered = route in routes_registered

        import time
        ip_batch = int(time.time()) % 200
        headers = {"X-Forwarded-For": f"172.16.{ip_batch}.{idx}"}

        try:
            if files_payload:
                if "files" in files_payload and isinstance(files_payload["files"], list):
                    f_list = files_payload["files"]
                else:
                    f_list = [(k, v) for k, v in files_payload.items()]
                res = client.post(route, files=f_list, data=data_payload, headers=headers)
            else:
                res = client.post(route, data=data_payload, headers=headers)

            success = res.status_code == 200
            content_length = len(res.content)
            output_valid = success and content_length > 0

            status = "PASS" if (is_registered and success and output_valid) else "FAIL"
            err_msg = "" if status == "PASS" else f"HTTP {res.status_code}: {res.text[:100]}"

            results.append({
                "num": idx,
                "name": name,
                "tool_key": tool_key,
                "old_route": old_route,
                "new_route": route,
                "registered": is_registered,
                "http_status": res.status_code,
                "processing_test": success,
                "output_valid": output_valid,
                "status": status,
                "error": err_msg
            })

            print(f"[{status:4}] {idx:2}. {name:20} | Route: {route:30} | HTTP {res.status_code} | Output: {content_length} bytes")

        except Exception as exc:
            results.append({
                "num": idx,
                "name": name,
                "tool_key": tool_key,
                "old_route": old_route,
                "new_route": route,
                "registered": is_registered,
                "http_status": 500,
                "processing_test": False,
                "output_valid": False,
                "status": "FAIL",
                "error": str(exc)
            })
            print(f"[FAIL] {idx:2}. {name:20} | Route: {route:30} | EXCEPTION: {exc}")

    passed_count = sum(1 for r in results if r["status"] == "PASS")
    failed_count = sum(1 for r in results if r["status"] == "FAIL")

    print(f"\n==================================================")
    print(f"FINAL RESULT: {passed_count}/48 PASSED ({failed_count} FAILED)")
    print(f"==================================================")

    return results

if __name__ == "__main__":
    run_v1_matrix_test()
